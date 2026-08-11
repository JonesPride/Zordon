from __future__ import annotations

import os
from pathlib import Path

import pytest
from docx import Document
from pypdf import PdfWriter

from zordon.config import ApprovedRoot, Tier2Limits
from zordon.documents import DocumentService
from zordon.paths import ApprovedFolders


def service_for(
    root: Path,
    *,
    folder_id: str = "docs",
    document_scan: int = 100,
    search_results: int = 10,
    document_read_chars: int = 12_000,
) -> DocumentService:
    root.mkdir(exist_ok=True)
    folders = ApprovedFolders((ApprovedRoot(folder_id, root.resolve()),))
    limits = Tier2Limits(
        document_scan=document_scan,
        search_results=search_results,
        document_read_chars=document_read_chars,
    )
    return DocumentService(folders, limits)


def test_text_accepts_utf8_bom_and_markdown_case_insensitively(tmp_path: Path) -> None:
    root = tmp_path / "root"
    root.mkdir()
    (root / "Notes.MD").write_bytes(b"\xef\xbb\xbf# Heading\n\nCaf\xc3\xa9")

    result = service_for(root).read("docs", "Notes.MD")

    assert result.ok is True
    assert result.data["text"] == "# Heading\n\nCaf\u00e9"
    assert result.data["format"] == "md"
    assert result.data["untrusted_data"] is True


def test_text_rejects_invalid_utf8_without_guessing(tmp_path: Path) -> None:
    root = tmp_path / "root"
    root.mkdir()
    (root / "bad.txt").write_bytes(b"hello\xff")

    result = service_for(root).read("docs", "bad.txt")

    assert result.ok is False
    assert result.code == "document_encoding_invalid"
    assert str(root) not in result.summary


def test_docx_extracts_paragraphs_then_table_cells(tmp_path: Path) -> None:
    root = tmp_path / "root"
    root.mkdir()
    path = root / "notes.docx"
    document = Document()
    document.add_paragraph("Project heading")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "ComicVine"
    table.cell(0, 1).text = "metadata"
    document.save(str(path))

    result = service_for(root).read("docs", "notes.docx")

    assert result.ok is True
    assert result.data["text"] == "Project heading\nComicVine\nmetadata"


def test_pdf_joins_extracted_page_text(monkeypatch: pytest.MonkeyPatch, tmp_path: Path) -> None:
    root = tmp_path / "root"
    root.mkdir()
    path = root / "report.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=72, height=72)
    writer.add_blank_page(width=72, height=72)
    with path.open("wb") as stream:
        writer.write(stream)

    from zordon import documents

    extracted = iter(("Page 1", "Page 2"))
    monkeypatch.setattr(documents, "_extract_pdf_page", lambda page: next(extracted))
    result = service_for(root).read("docs", "report.pdf")

    assert result.ok is True
    assert result.data["text"] == "Page 1\nPage 2"


def test_pdf_rejects_encrypted_and_more_than_300_pages(tmp_path: Path) -> None:
    root = tmp_path / "root"
    root.mkdir()
    encrypted = root / "encrypted.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=72, height=72)
    writer.encrypt("secret")
    with encrypted.open("wb") as stream:
        writer.write(stream)
    too_long = root / "long.pdf"
    writer = PdfWriter()
    for _ in range(301):
        writer.add_blank_page(width=1, height=1)
    with too_long.open("wb") as stream:
        writer.write(stream)

    service = service_for(root)

    assert service.read("docs", "encrypted.pdf").code == "document_encrypted"
    assert service.read("docs", "long.pdf").code == "document_page_limit"


@pytest.mark.parametrize(("name", "content"), [("bad.pdf", b"not pdf"), ("bad.docx", b"not zip")])
def test_malformed_binary_document_returns_stable_safe_error(
    tmp_path: Path, name: str, content: bytes
) -> None:
    root = tmp_path / "root"
    root.mkdir()
    (root / name).write_bytes(content)

    result = service_for(root).read("docs", name)

    assert result.ok is False
    assert result.code == "document_malformed"
    assert str(root) not in result.summary
    assert result.data == {}


def test_unexpected_parser_failure_is_wrapped_without_exception_text(
    monkeypatch: pytest.MonkeyPatch, tmp_path: Path
) -> None:
    root = tmp_path / "root"
    root.mkdir()
    path = root / "report.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=72, height=72)
    with path.open("wb") as stream:
        writer.write(stream)
    from zordon import documents

    monkeypatch.setattr(
        documents,
        "_extract_pdf_page",
        lambda page: (_ for _ in ()).throw(RuntimeError("private parser detail")),
    )

    result = service_for(root).read("docs", "report.pdf")

    assert result.code == "document_extraction_failed"
    assert "private parser detail" not in result.summary


@pytest.mark.parametrize(
    ("name", "size"),
    [("huge.txt", 2 * 1024 * 1024 + 1), ("huge.pdf", 20 * 1024 * 1024 + 1)],
)
def test_source_size_caps_are_checked_before_parsing(
    tmp_path: Path, name: str, size: int
) -> None:
    root = tmp_path / "root"
    root.mkdir()
    with (root / name).open("wb") as stream:
        stream.truncate(size)

    result = service_for(root).read("docs", name)

    assert result.code == "document_too_large"


def test_extraction_over_one_million_characters_is_rejected(
    monkeypatch: pytest.MonkeyPatch, tmp_path: Path
) -> None:
    root = tmp_path / "root"
    root.mkdir()
    path = root / "large.docx"
    document = Document()
    document.add_paragraph("small")
    document.save(str(path))
    from zordon import documents

    monkeypatch.setattr(documents, "_extract_docx", lambda _: "x" * 1_000_001)

    assert service_for(root).read("docs", "large.docx").code == "document_text_limit"


def test_cache_reuses_matching_stat_and_invalidates_after_change(
    monkeypatch: pytest.MonkeyPatch, tmp_path: Path
) -> None:
    root = tmp_path / "root"
    root.mkdir()
    path = root / "note.txt"
    path.write_text("first")
    from zordon import documents

    original = documents._extract_text
    calls = 0

    def counted(candidate: Path) -> str:
        nonlocal calls
        calls += 1
        return original(candidate)

    monkeypatch.setattr(documents, "_extract_text", counted)
    service = service_for(root)
    assert service.read("docs", "note.txt").data["text"] == "first"
    assert service.read("docs", "note.txt").data["text"] == "first"
    path.write_text("second")
    stat = path.stat()
    os.utime(path, ns=(stat.st_atime_ns, stat.st_mtime_ns + 1_000_000))
    assert service.read("docs", "note.txt").data["text"] == "second"
    assert calls == 2


def test_read_returns_bounded_window_and_identity(tmp_path: Path) -> None:
    root = tmp_path / "root"
    root.mkdir()
    (root / "note.txt").write_text("0123456789")

    result = service_for(root, document_read_chars=1_000).read(
        "docs", "note.txt", start_char=3, max_chars=4
    )

    assert result.data == {
        "folder_id": "docs",
        "relative_path": "note.txt",
        "format": "txt",
        "text": "3456",
        "start_char": 3,
        "end_char": 7,
        "total_chars": 10,
        "truncated": True,
        "untrusted_data": True,
    }


def test_read_past_end_returns_a_coherent_empty_range(tmp_path: Path) -> None:
    root = tmp_path / "root"
    root.mkdir()
    (root / "note.txt").write_text("short")

    result = service_for(root).read("docs", "note.txt", start_char=100, max_chars=4)

    assert result.data["text"] == ""
    assert result.data["start_char"] == 5
    assert result.data["end_char"] == 5
    assert result.data["truncated"] is True


@pytest.mark.parametrize(("start", "maximum"), [(-1, 5), (0, 0), (0, 1_001)])
def test_read_rejects_invalid_windows(
    tmp_path: Path, start: int, maximum: int
) -> None:
    root = tmp_path / "root"
    root.mkdir()
    (root / "note.txt").write_text("text")

    result = service_for(root, document_read_chars=1_000).read(
        "docs", "note.txt", start_char=start, max_chars=maximum
    )

    assert result.code == "document_read_invalid"


def test_read_rejects_unsupported_format_and_safe_path_errors(tmp_path: Path) -> None:
    root = tmp_path / "root"
    root.mkdir()
    (root / "data.csv").write_text("a,b")
    service = service_for(root)

    assert service.read("docs", "data.csv").code == "document_format_unsupported"
    missing = service.read("docs", "missing.txt")
    assert missing.code == "file_not_found"
    assert str(root) not in missing.summary


def test_search_ranks_phrase_filename_heading_then_body(tmp_path: Path) -> None:
    root = tmp_path / "root"
    root.mkdir()
    (root / "z-phrase.txt").write_text("Report\nThe alpha beta phrase is exact.")
    (root / "alpha beta filename.txt").write_text("unrelated")
    (root / "heading.md").write_text("# Alpha beta\nother")
    (root / "body.txt").write_text("alpha appears far from beta")

    result = service_for(root).search("alpha beta")

    assert result.ok is True
    assert [item["relative_path"] for item in result.data["results"]] == [
        "z-phrase.txt",
        "alpha beta filename.txt",
        "heading.md",
        "body.txt",
    ]
    assert result.data["untrusted_data"] is True
    assert all(len(item["snippet"]) <= 500 for item in result.data["results"])


def test_search_filters_folders_and_breaks_ties_deterministically(tmp_path: Path) -> None:
    alpha = tmp_path / "alpha"
    beta = tmp_path / "beta"
    alpha.mkdir()
    beta.mkdir()
    (alpha / "b.txt").write_text("needle")
    (alpha / "a.txt").write_text("needle")
    (beta / "a.txt").write_text("needle")
    folders = ApprovedFolders(
        (ApprovedRoot("beta", beta.resolve()), ApprovedRoot("alpha", alpha.resolve()))
    )
    service = DocumentService(folders, Tier2Limits(document_scan=100))

    all_results = service.search("needle")
    filtered = service.search("needle", folder_ids=["BETA"])

    assert [(item["folder_id"], item["relative_path"]) for item in all_results.data["results"]] == [
        ("alpha", "a.txt"),
        ("alpha", "b.txt"),
        ("beta", "a.txt"),
    ]
    assert [item["folder_id"] for item in filtered.data["results"]] == ["beta"]


def test_search_obeys_requested_and_configured_result_limits(tmp_path: Path) -> None:
    root = tmp_path / "root"
    root.mkdir()
    for index in range(5):
        (root / f"{index}.txt").write_text("needle")
    service = service_for(root, search_results=3)

    assert len(service.search("needle").data["results"]) == 3
    assert len(service.search("needle", max_results=2).data["results"]) == 2
    assert service.search("needle", max_results=4).code == "document_search_invalid"


def test_search_without_approved_roots_is_controlled() -> None:
    service = DocumentService(ApprovedFolders(()), Tier2Limits())

    result = service.search("needle")

    assert result.ok is False
    assert result.code == "approved_folders_missing"


def test_search_skips_bad_file_and_reports_warning(tmp_path: Path) -> None:
    root = tmp_path / "root"
    root.mkdir()
    (root / "good.txt").write_text("needle")
    (root / "bad.txt").write_bytes(b"\xff")

    result = service_for(root).search("needle")

    assert [item["relative_path"] for item in result.data["results"]] == ["good.txt"]
    assert result.data["warning_count"] == 1


def test_search_marks_partial_when_scan_limit_is_hit(tmp_path: Path) -> None:
    root = tmp_path / "root"
    root.mkdir()
    (root / "a.txt").write_text("needle")
    (root / "b.txt").write_text("needle")

    result = service_for(root, document_scan=1).search("needle")

    assert result.data["partial"] is True
    assert len(result.data["results"]) == 1


@pytest.mark.parametrize("query", ["", "   "])
def test_search_rejects_blank_query(tmp_path: Path, query: str) -> None:
    root = tmp_path / "root"
    root.mkdir()
    assert service_for(root).search(query).code == "document_search_invalid"
