from __future__ import annotations

from collections.abc import Iterable
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from docx import Document
from pypdf import PageObject, PdfReader
from pypdf.errors import PdfReadError

from zordon.config import Tier2Limits
from zordon.paths import ApprovedFolderError, ApprovedFolders, ResolvedFile
from zordon.tools.base import ToolResult

_SUPPORTED_SUFFIXES = {".txt", ".md", ".pdf", ".docx"}
_TEXT_SOURCE_LIMIT = 2 * 1024 * 1024
_BINARY_SOURCE_LIMIT = 20 * 1024 * 1024
_PDF_PAGE_LIMIT = 300
_EXTRACTED_TEXT_LIMIT = 1_000_000
_SNIPPET_LIMIT = 500


class DocumentError(RuntimeError):
    def __init__(self, code: str, summary: str) -> None:
        super().__init__(summary)
        self.code = code


@dataclass(frozen=True, slots=True)
class _ExtractedDocument:
    file: ResolvedFile
    text: str
    format: str


class DocumentService:
    """Search and read bounded text from approved local documents."""

    def __init__(self, folders: ApprovedFolders, limits: Tier2Limits) -> None:
        self._folders = folders
        self._limits = limits
        self._cache: dict[tuple[Path, int, int], str] = {}

    def read(
        self,
        folder_id: str,
        relative_path: str,
        start_char: int = 0,
        max_chars: int | None = None,
    ) -> ToolResult:
        window = self._limits.document_read_chars if max_chars is None else max_chars
        if (
            isinstance(start_char, bool)
            or not isinstance(start_char, int)
            or start_char < 0
            or isinstance(window, bool)
            or not isinstance(window, int)
            or not 1 <= window <= self._limits.document_read_chars
        ):
            return ToolResult.failure(
                "document_read_invalid",
                "The document read window is outside the configured limits.",
            )
        try:
            file = self._folders.resolve_file(folder_id, relative_path)
            document = self._extract(file)
        except (ApprovedFolderError, DocumentError) as exc:
            return ToolResult.failure(exc.code, str(exc))

        total = len(document.text)
        actual_start = min(start_char, total)
        end = min(actual_start + window, total)
        text = document.text[actual_start:end]
        return ToolResult.success(
            "Document text read.",
            {
                "folder_id": file.folder_id,
                "relative_path": file.relative_path,
                "format": document.format,
                "text": text,
                "start_char": actual_start,
                "end_char": end,
                "total_chars": total,
                "truncated": start_char > 0 or end < total,
                "untrusted_data": True,
            },
        )

    def search(
        self,
        query: str,
        folder_ids: Iterable[str] | None = None,
        max_results: int | None = None,
    ) -> ToolResult:
        normalized_query = query.strip().casefold() if isinstance(query, str) else ""
        result_limit = self._limits.search_results if max_results is None else max_results
        if (
            not normalized_query
            or isinstance(result_limit, bool)
            or not isinstance(result_limit, int)
            or not 1 <= result_limit <= self._limits.search_results
        ):
            return ToolResult.failure(
                "document_search_invalid",
                "The document search request is invalid.",
            )

        selected = tuple(folder_ids) if folder_ids is not None else ()
        if not selected:
            selected = self._folders.folder_ids
        if not selected:
            return ToolResult.failure(
                "approved_folders_missing",
                "No approved folders are configured for document search.",
            )

        try:
            scanned = self._folders.iter_files(
                selected,
                _SUPPORTED_SUFFIXES,
                self._limits.document_scan + 1,
            )
        except ApprovedFolderError as exc:
            return ToolResult.failure(exc.code, str(exc))

        partial = len(scanned) > self._limits.document_scan
        files = scanned[: self._limits.document_scan]
        tokens = tuple(dict.fromkeys(normalized_query.split()))
        ranked: list[tuple[float, str, str, dict[str, Any]]] = []
        warning_count = 0
        for file in files:
            try:
                document = self._extract(file)
            except DocumentError:
                warning_count += 1
                continue
            score = _score(document, normalized_query, tokens)
            if score <= 0:
                continue
            ranked.append(
                (
                    -score,
                    file.folder_id.casefold(),
                    file.relative_path.casefold(),
                    {
                        "folder_id": file.folder_id,
                        "relative_path": file.relative_path,
                        "format": document.format,
                        "snippet": _snippet(document.text, normalized_query, tokens),
                        "score": score,
                    },
                )
            )
        ranked.sort(key=lambda item: item[:3])
        return ToolResult.success(
            "Document search complete.",
            {
                "query": query.strip(),
                "results": [item[3] for item in ranked[:result_limit]],
                "partial": partial,
                "warning_count": warning_count,
                "untrusted_data": True,
            },
        )

    def _extract(self, file: ResolvedFile) -> _ExtractedDocument:
        suffix = file.path.suffix.casefold()
        if suffix not in _SUPPORTED_SUFFIXES:
            raise DocumentError(
                "document_format_unsupported",
                "The requested file is not a supported document format.",
            )
        source_limit = (
            _TEXT_SOURCE_LIMIT if suffix in {".txt", ".md"} else _BINARY_SOURCE_LIMIT
        )
        if file.size > source_limit:
            raise DocumentError(
                "document_too_large",
                "The document exceeds the configured source-size limit.",
            )

        key = (file.path, file.size, file.mtime_ns)
        text = self._cache.get(key)
        if text is None:
            text = _extract_file(file.path, suffix)
            if len(text) > _EXTRACTED_TEXT_LIMIT:
                raise DocumentError(
                    "document_text_limit",
                    "The document exceeds the extracted-text limit.",
                )
            stale = [cache_key for cache_key in self._cache if cache_key[0] == file.path]
            for cache_key in stale:
                del self._cache[cache_key]
            self._cache[key] = text
        return _ExtractedDocument(file, text, suffix.removeprefix("."))


def _extract_file(path: Path, suffix: str) -> str:
    try:
        if suffix in {".txt", ".md"}:
            return _extract_text(path)
        if suffix == ".pdf":
            return _extract_pdf(path)
        return _extract_docx(path)
    except DocumentError:
        raise
    except UnicodeError as exc:
        raise DocumentError(
            "document_encoding_invalid",
            "The document is not valid UTF-8 text.",
        ) from exc
    except Exception as exc:
        raise DocumentError(
            "document_extraction_failed",
            "The document could not be extracted safely.",
        ) from exc


def _extract_text(path: Path) -> str:
    return path.read_text(encoding="utf-8-sig", errors="strict")


def _extract_pdf(path: Path) -> str:
    try:
        reader = PdfReader(path)
        if reader.is_encrypted:
            raise DocumentError(
                "document_encrypted",
                "Encrypted documents cannot be read.",
            )
        if len(reader.pages) > _PDF_PAGE_LIMIT:
            raise DocumentError(
                "document_page_limit",
                "The PDF exceeds the page limit.",
            )
        return "\n".join(_extract_pdf_page(page) or "" for page in reader.pages)
    except DocumentError:
        raise
    except (OSError, PdfReadError, ValueError) as exc:
        raise DocumentError(
            "document_malformed",
            "The document could not be parsed safely.",
        ) from exc


def _extract_pdf_page(page: PageObject) -> str:
    return page.extract_text() or ""


def _extract_docx(path: Path) -> str:
    try:
        document = Document(str(path))
        parts = [paragraph.text for paragraph in document.paragraphs]
        parts.extend(
            cell.text
            for table in document.tables
            for row in table.rows
            for cell in row.cells
        )
        return "\n".join(parts)
    except (OSError, ValueError, KeyError) as exc:
        raise DocumentError(
            "document_malformed",
            "The document could not be parsed safely.",
        ) from exc
    except Exception as exc:
        if exc.__class__.__module__.startswith(("docx", "zipfile")):
            raise DocumentError(
                "document_malformed",
                "The document could not be parsed safely.",
            ) from exc
        raise


def _score(
    document: _ExtractedDocument,
    phrase: str,
    tokens: tuple[str, ...],
) -> float:
    body = document.text.casefold()
    filename = Path(document.file.relative_path).stem.casefold()
    lines = body.splitlines()
    title_index = next((index for index, line in enumerate(lines) if line.strip()), None)
    first_line = lines[title_index].strip() if title_index is not None else ""
    searchable_body = "\n".join(
        line for index, line in enumerate(lines) if index != title_index
    )
    if phrase in searchable_body:
        return 100.0
    if phrase in filename:
        return 60.0
    if phrase in first_line:
        return 40.0
    matched = sum(token in body for token in tokens)
    return 20.0 * matched / len(tokens) if tokens else 0.0


def _snippet(text: str, phrase: str, tokens: tuple[str, ...]) -> str:
    folded = text.casefold()
    match = folded.find(phrase)
    if match < 0:
        positions = [folded.find(token) for token in tokens]
        match = min((position for position in positions if position >= 0), default=0)
    start = max(0, match - _SNIPPET_LIMIT // 2)
    end = min(len(text), start + _SNIPPET_LIMIT)
    start = max(0, end - _SNIPPET_LIMIT)
    return text[start:end]
